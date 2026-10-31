"""
Daftar unit PO yang BELUM masuk invoice (Histori Dokumen) untuk satu pelanggan, siap jadi Excel.
Patokan "sudah ditagih" = unit tercantum di invoice yang tersimpan di Histori Dokumen:
  1) No. Rangka unit muncul di baris invoice (kolom Keterangan), atau
  2) unit tanpa No. Rangka, tapi order_id-nya tercantum di order_ids sebuah invoice.
Flag status_invoice di PO tidak dipakai sebagai patokan (cuma ditampilkan sebagai pembanding).
"""
import io
import re

_BADAN = re.compile(r"\b(PT|CV|UD|TBK)\b\.?", re.I)


def norm_nama(s):
    s = _BADAN.sub(" ", (s or "").upper())
    return re.sub(r"[^A-Z0-9]+", " ", s).strip()


def cocok_nama(nama, query):
    q = norm_nama(query)
    return bool(q) and q in norm_nama(nama)


def _rangka_key(s):
    return re.sub(r"[^A-Z0-9]", "", (s or "").upper())


def hitung(orders, invoices, query):
    """orders: dokumen PO; invoices: dokumen doc_history jenis invoice. Mengembalikan (baris, ringkas)."""
    teks_inv, oid_inv = [], {}
    n_inv = 0
    for d in invoices:
        if d.get("jenis") not in (None, "invoice"):
            continue
        if not (cocok_nama(d.get("customer"), query) or cocok_nama((d.get("meta") or {}).get("customer_nama"), query)):
            continue
        n_inv += 1
        no = d.get("no_dokumen") or d.get("id") or "(tanpa nomor)"
        for o in d.get("order_ids") or []:
            oid_inv.setdefault(str(o), no)
        for l in d.get("lines") or []:
            teks_inv.append((_rangka_key(re.sub(r"<[^>]+>", " ", str(l.get("ket") or ""))), no))

    def cari_inv(rk):
        for t, no in teks_inv:
            if rk in t:
                return no
        return ""

    baris, sudah, nama_ct, sudah_list = [], 0, {}, []
    for o in orders:
        if (o.get("status") or "").upper() == "CANCELLED" or not cocok_nama(o.get("customer_nama"), query):
            continue
        nm = (o.get("customer_nama") or "").strip()
        for u in (o.get("units") or [{}]):
            nama_ct[nm] = nama_ct.get(nm, 0) + 1
            rk = _rangka_key(u.get("no_rangka") or o.get("no_rangka"))
            no_inv = (cari_inv(rk) if len(rk) >= 6 else "") if rk else oid_inv.get(str(o.get("order_id")), "")
            if no_inv:
                sudah += 1
                sudah_list.append({"order_id": o.get("order_id", ""), "tanggal": (o.get("created_at") or "")[:10],
                                   "no_rangka": u.get("no_rangka") or o.get("no_rangka", ""),
                                   "tujuan": u.get("tujuan") or o.get("tujuan_kota", ""), "invoice": no_inv})
                continue
            baris.append({
                "order_id": o.get("order_id", ""),
                "tanggal": (o.get("created_at") or "")[:10],
                "asal": o.get("asal_kota", ""), "tujuan": u.get("tujuan") or o.get("tujuan_kota", ""),
                "kendaraan": " ".join(x for x in [(u.get("vehicle_type") or o.get("vehicle_type") or ""), (u.get("tipe_model") or "")] if x),
                "nopol": u.get("nopol") or o.get("nopol", ""),
                "no_rangka": u.get("no_rangka") or o.get("no_rangka", ""),
                "no_mesin": u.get("no_mesin", ""), "warna": u.get("warna") or o.get("warna", ""),
                "kapal": u.get("nama_kapal", ""), "etd": u.get("etd", ""), "serah_terima": u.get("serah_terima", ""),
                "status_po": o.get("status", ""), "status_perjalanan": u.get("status_perjalanan", ""),
                "flag_invoice": u.get("status_invoice", ""), "trip_id": o.get("trip_id") or "",
            })
    baris.sort(key=lambda r: (r["tanggal"], r["order_id"]))
    return baris, {"invoice": n_inv, "unit_sudah_ditagih": sudah, "unit_belum_ditagih": len(baris),
                   "nama": sorted(nama_ct.items(), key=lambda x: -x[1]), "sudah": sudah_list}


KOLOM = [
    ("No. PO (Order)", "order_id", 16), ("Tgl PO", "tanggal", 12), ("Asal", "asal", 16), ("Tujuan", "tujuan", 18),
    ("Kendaraan", "kendaraan", 24), ("Nopol", "nopol", 12), ("No. Rangka", "no_rangka", 22), ("No. Mesin", "no_mesin", 16),
    ("Warna", "warna", 10), ("Kapal", "kapal", 16), ("ETD", "etd", 12), ("Serah Terima", "serah_terima", 13),
    ("Status PO", "status_po", 12), ("Status Perjalanan", "status_perjalanan", 16),
    ("Flag Invoice di PO", "flag_invoice", 16), ("Trip ID", "trip_id", 16),
]


def ke_xlsx(baris, judul, ring=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Belum Ditagih"
    ws["A1"] = judul
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["No"] + [k[0] for k in KOLOM])
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, r in enumerate(baris, 1):
        ws.append([i] + [r.get(k[1], "") for k in KOLOM])
    ws.column_dimensions["A"].width = 5
    for j, k in enumerate(KOLOM, 2):
        ws.column_dimensions[ws.cell(row=3, column=j).column_letter].width = k[2]
    ws.freeze_panes = "A4"
    ws.append([])
    ws.append(["", f"Total {len(baris)} unit belum ditagih"])
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
    if ring:
        w2 = wb.create_sheet("Sudah Ditagih")
        w2.append(["No. PO", "Tgl PO", "No. Rangka", "Tujuan", "Ada di Invoice"])
        for c in w2[1]:
            c.font = Font(bold=True)
        for r in ring.get("sudah", []):
            w2.append([r["order_id"], r["tanggal"], r["no_rangka"], r["tujuan"], r["invoice"]])
        for col, wd in zip("ABCDE", (16, 12, 22, 18, 22)):
            w2.column_dimensions[col].width = wd
        w3 = wb.create_sheet("Nama Pelanggan")
        w3.append(["Nama pelanggan di PO yang ikut terbaca", "Jumlah unit"])
        for c in w3[1]:
            c.font = Font(bold=True)
        for nm, n in ring.get("nama", []):
            w3.append([nm, n])
        w3.column_dimensions["A"].width = 48
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def peta_status(orders, invoices):
    """Status tagih per PO (semua pelanggan): {order_id: "selesai" | "belum_tagih"}.
    "selesai" = SEMUA unit PO sudah tercantum di invoice Histori Dokumen (patokan sama dengan `hitung`);
    "belum_tagih" = masih ada unit yang belum. PO Cancelled tidak diberi status."""
    teks, oid_inv = [], set()
    for d in invoices:
        if d.get("jenis") not in (None, "invoice"):
            continue
        for o in d.get("order_ids") or []:
            oid_inv.add(str(o))
        for l in d.get("lines") or []:
            teks.append(_rangka_key(re.sub(r"<[^>]+>", " ", str(l.get("ket") or ""))))
    gabung = " ".join(teks)
    out = {}
    for o in orders:
        if (o.get("status") or "").upper() == "CANCELLED" or not o.get("order_id"):
            continue
        semua = True
        for u in (o.get("units") or [{}]):
            rk = _rangka_key(u.get("no_rangka") or o.get("no_rangka"))
            ok = (len(rk) >= 6 and rk in gabung) if rk else str(o["order_id"]) in oid_inv
            if not ok:
                semua = False
                break
        out[o["order_id"]] = "selesai" if semua else "belum_tagih"
    return out
